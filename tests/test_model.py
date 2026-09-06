import hashlib
import json
import os
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd
import pytest

from src.pipeline import run_pipeline

ROOT = Path(__file__).resolve().parents[1]


def test_grains_bookings_and_revenue_from_source(pipeline_output):
    output, _ = pipeline_output
    workbook = openpyxl.load_workbook(ROOT / 'data/raw/UseCase_Airlines.xlsx', read_only=True, data_only=True)
    payment_rows = list(workbook['payments'].values)
    amount_index = payment_rows[0].index('amount')
    expected = sum((Decimal(str(row[amount_index])) for row in payment_rows[1:]
                    if isinstance(row[amount_index], (int, float))), Decimal(0))
    workbook.close()
    with duckdb.connect(str(output / 'data/gold/asg_airlines.duckdb'), read_only=True) as con:
        assert con.execute('SELECT count(*), count(DISTINCT flight_sk) FROM fact_flight').fetchone() == (1004, 1004)
        assert con.execute('SELECT count(*) FROM fact_flight WHERE NOT is_unknown_member').fetchone()[0] == 1003
        assert con.execute('SELECT count(*), count(DISTINCT booking_id) FROM fact_booking').fetchone() == (1000, 1000)
        assert con.execute('SELECT count(*), count(DISTINCT payment_id) FROM fact_payment').fetchone() == (1000, 1000)
        assert con.execute("SELECT sum(amount) FROM fact_payment WHERE amount_quality='valid'").fetchone()[0] == expected == Decimal('7385142.98')
        assert con.execute('SELECT sum(p.amount) FROM fact_payment p JOIN fact_booking b USING (booking_sk) JOIN fact_flight f USING (flight_sk)').fetchone()[0] == expected
        assert con.execute('SELECT count(*) FROM fact_booking WHERE flight_sk=-1').fetchone()[0] == 2
        assert con.execute("SELECT count(*) FROM fact_flight WHERE flight_id='6F250'").fetchone()[0] == 0
        assert con.execute('SELECT sum(gross_valid_payment_amount) FROM kpi_route').fetchone()[0] == expected


def test_required_kpi_denominators(pipeline_output):
    _, manifest = pipeline_output
    kpi = manifest['kpis']
    assert kpi['cancellation_rate'] == pytest.approx(314 / 1000)
    assert kpi['confirmation_rate'] == pytest.approx(320 / 1000)
    assert kpi['payment_coverage'] == pytest.approx(637 / 1000)
    assert kpi['overnight_share'] == pytest.approx(122 / 1003)
    assert kpi['red_eye_share'] == pytest.approx(271 / 1003)
    assert kpi['bookings_without_payment'] == 363


def test_csv_types_and_unknown_flight(pipeline_output):
    output, manifest = pipeline_output
    for name, count in manifest['gold_rows'].items():
        frame = pd.read_csv(output / 'data/gold' / f'{name}.csv', keep_default_na=False)
        assert len(frame) == count
    flights = pd.read_csv(output / 'data/gold/fact_flight.csv')
    assert flights.loc[flights.flight_sk == -1, 'duration_minutes'].isna().all()
    assert flights.loc[flights.flight_id == 'SJ192', 'duration_minutes'].item() == 300
    assert not flights.loc[flights.flight_id == 'SJ192', 'is_overnight'].item()


def test_rerun_same_input_same_analytical_csvs(pipeline_output):
    output, first = pipeline_output
    second = run_pipeline(ROOT / 'data/raw/UseCase_Airlines.xlsx', output,
                          pepper='unit-test-pepper-for-synthetic-checks-only')
    assert second['run_id'] != first['run_id']
    assert second['gold_sha256'] == first['gold_sha256']
    assert second['kpis'] == first['kpis']


def test_failure_preserves_last_success_and_reports_stage(pipeline_output, tmp_path):
    output, _ = pipeline_output
    before = (output / 'reports/run_manifest.json').read_bytes()
    csv_hash = hashlib.sha256((output / 'data/gold/fact_payment.csv').read_bytes()).hexdigest()
    workbook = openpyxl.Workbook()
    workbook.active.title = 'flights'
    workbook.active.append(['unexpected_column'])
    workbook.active.append(['bad'])
    malformed = tmp_path / 'missing_sheets.xlsx'
    workbook.save(malformed)
    with pytest.raises(RuntimeError, match='validate'):
        run_pipeline(malformed, output, pepper='unit-test-pepper-for-synthetic-checks-only')
    assert (output / 'reports/run_manifest.json').read_bytes() == before
    assert hashlib.sha256((output / 'data/gold/fact_payment.csv').read_bytes()).hexdigest() == csv_hash
    failure = json.loads(max((output / 'reports/failed_runs').glob('*.json')).read_text())
    assert failure['status'] == 'failed'
    assert failure['stages'][-1]['name'] == 'validate'
    assert failure['stages'][-1]['status'] == 'failed'


def test_cli_without_pepper_exits_nonzero(tmp_path):
    env = os.environ.copy()
    env.pop('ASG_PII_PEPPER', None)
    result = subprocess.run([sys.executable, '-m', 'src.pipeline', '--output-dir', str(tmp_path)],
                            cwd=ROOT, env=env, capture_output=True, text=True, check=False)
    assert result.returncode != 0
    assert 'ASG_PII_PEPPER' in result.stderr
    assert not (tmp_path / 'data/gold').exists()


def test_corrupt_workbook_fixture_fails_ingestion_without_publishing(tmp_path):
    fixture = ROOT / 'data/raw/malformed_corrupt.xlsx'
    if not fixture.exists():
        # Raw files are excluded from Git; keep this failure check usable in a checkout.
        fixture = tmp_path / 'malformed_corrupt.xlsx'
        fixture.write_bytes(b'This is deliberately not an XLSX ZIP archive.')
    output = tmp_path / 'result'
    with pytest.raises(RuntimeError, match=r'ingest \(BadZipFile\)'):
        run_pipeline(fixture, output, pepper='unit-test-pepper-for-synthetic-checks-only')
    assert not (output / 'data/gold').exists()
    assert not (output / 'reports/run_manifest.json').exists()
    failures = list((output / 'reports/failed_runs').glob('*.json'))
    assert len(failures) == 1
    failure = json.loads(failures[0].read_text())
    assert failure['status'] == 'failed'
    assert failure['error_type'] == 'BadZipFile'
    assert failure['stages'][-1]['name'] == 'ingest'
    assert failure['stages'][-1]['status'] == 'failed'


def test_bad_row_and_uncached_durations_do_not_fail_batch(tmp_path):
    workbook = openpyxl.load_workbook(ROOT / 'data/raw/UseCase_Airlines.xlsx')
    flights = workbook['flights']
    headers = [cell.value for cell in flights[1]]
    changed_flight = flights.cell(2, headers.index('flight_id') + 1).value
    flights.cell(2, headers.index('arrival_time') + 1, 'not-a-timestamp')
    bookings = workbook['bookings']
    booking_headers = [cell.value for cell in bookings[1]]
    changed_booking = bookings.cell(2, booking_headers.index('booking_id') + 1).value
    bookings.cell(2, booking_headers.index('status') + 1, 'UNSUPPORTED')
    mutated = tmp_path / 'dirty.xlsx'
    workbook.save(mutated)
    workbook.close()
    output = tmp_path / 'result'
    manifest = run_pipeline(mutated, output, pepper='unit-test-pepper-for-synthetic-checks-only')
    assert manifest['status'] == 'success'
    assert manifest['silver_rows']['bookings'] == 1000
    assert manifest['quarantine_rows']['flights'] >= 3
    with duckdb.connect(str(output / 'data/gold/asg_airlines.duckdb'), read_only=True) as con:
        assert con.execute("SELECT s.status FROM fact_booking b JOIN dim_status s USING(status_sk) WHERE b.booking_id=?", [changed_booking]).fetchone()[0] == 'INVALID'
        assert con.execute('SELECT count(*) FROM fact_flight WHERE flight_id=?', [changed_flight]).fetchone()[0] == 0
        assert con.execute('SELECT count(*) FROM fact_flight WHERE duration_minutes IS NOT NULL').fetchone()[0] == manifest['silver_rows']['flights']


def test_decimal_warehouse_amount_boundary():
    from src.validate import classify_amount
    assert classify_amount('9999999999999999.99') == ('valid', 999999999999999999)
    assert classify_amount('10000000000000000.00') == ('out_of_range', None)


def test_booking_month_extract_matches_source_without_payment_fanout(pipeline_output):
    from collections import Counter, defaultdict

    output, manifest = pipeline_output
    workbook = openpyxl.load_workbook(ROOT / 'data/raw/UseCase_Airlines.xlsx', read_only=True, data_only=True)
    bookings = {row[0]: row for row in list(workbook['bookings'].values)[1:] if row[0] is not None}
    counts = Counter((row[3].year, row[3].month) for row in bookings.values())
    gross = defaultdict(Decimal)
    confirmed = defaultdict(Decimal)
    for row in list(workbook['payments'].values)[1:]:
        if isinstance(row[2], (int, float)):
            booking = bookings[row[1]]
            period = (booking[3].year, booking[3].month)
            gross[period] += Decimal(str(row[2]))
            if booking[4] == 'CONFIRMED':
                confirmed[period] += Decimal(str(row[2]))
    workbook.close()
    monthly = pd.read_csv(output / 'data/gold/kpi_booking_month.csv', dtype=str)
    assert monthly.columns.tolist() == ['year', 'month', 'month_label', 'bookings',
                                        'gross_valid_payment_amount', 'confirmed_payment_amount']
    assert manifest['gold_rows']['kpi_booking_month'] == len(monthly) == 13
    assert monthly.month_label.tolist() == pd.period_range('2025-04', '2026-04', freq='M').astype(str).tolist()
    for row in monthly.itertuples():
        period = (int(row.year), int(row.month))
        assert int(row.bookings) == counts[period]
        assert Decimal(row.gross_valid_payment_amount) == gross[period]
        assert Decimal(row.confirmed_payment_amount) == confirmed[period]
    assert monthly.bookings.astype(int).sum() == 1000
