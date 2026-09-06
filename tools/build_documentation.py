"""Build the case-study Word report and diagrams from a completed local run."""

import argparse
import csv
import json
from decimal import Decimal
from pathlib import Path

import matplotlib
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch

ROOT = Path(__file__).resolve().parents[1]
BLUE, INK, MUTED = "2E74B5", "17324D", "596777"


def draw_diagrams(folder):
    def canvas(height):
        fig, ax = plt.subplots(figsize=(10, height))
        fig.patch.set_facecolor("white")
        ax.set(xlim=(-.1, 10.1), ylim=(-.25, height+.25))
        ax.axis("off")
        return fig, ax

    def box(ax, x, y, label, width=2.02, height=.65, pending=False):
        ax.add_patch(FancyBboxPatch((x-width/2, y-height/2), width, height,
                     boxstyle="round,pad=0.055,rounding_size=0.08", linewidth=1.1,
                     edgecolor="#8B7553" if pending else "#AABCCB",
                     facecolor="#FFF8EB" if pending else "#EEF4F8"))
        ax.text(x, y, label, ha="center", va="center", fontsize=10.3, color="#17324D", linespacing=1.35)

    def arrow(ax, start, end, label=None, dashed=False):
        ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle": "->", "lw": 1.2,
                    "color": "#64798B", "linestyle": "--" if dashed else "-"})
        if label:
            ax.text((start[0]+end[0])/2, (start[1]+end[1])/2+.10, label,
                    fontsize=8.5, ha="center", color="#596777", bbox={"facecolor": "white", "edgecolor": "none", "pad": 1.2})

    def save(fig, name):
        fig.savefig(folder / f"{name}.png", dpi=220, bbox_inches="tight", pad_inches=.12)
        plt.close(fig)

    # The architecture PNG is rendered from architecture.mmd, not redrawn here.

    fig, ax = canvas(5.5)
    box(ax,5,5.05,"Excel sheets + source row trace",width=4.2)
    box(ax,5,3.92,"Validate structure and record row issues",width=4.7)
    arrow(ax,(5,4.68),(5,4.29))
    box(ax,1.15,3.92,"Schema failure\nstop + log",width=1.9,pending=True)
    arrow(ax,(2.57,3.92),(2.16,3.92))
    labels=["Flights\ntime + duplicates","Passengers\nsurvivor + token","Bookings\nstatus + references","Payments\namount quality"]
    for x,label in zip([1.2,3.7,6.2,8.7],labels):
        box(ax,x,2.6,label)
        arrow(ax,(5,3.55),(x,3.00))
        arrow(ax,(x,2.22),(5,1.57))
    box(ax,5,1.18,"Silver → unique keys → payment-grain facts",width=5.4)
    box(ax,5,.18,"Gold CSVs / DuckDB / quality evidence",width=5.4,height=.5)
    arrow(ax,(5,.8),(5,.5))
    ax.text(.15,.9,"Ambiguous or unusable\nrows go to restricted\nquarantine.",fontsize=9.1,color="#596777",ha="left")
    save(fig,"data_flow")

    fig, ax = canvas(5.75)
    for x,label in [(1.3,"dim_airline\nairline_sk"),(5,"dim_route\nroute_sk"),(8.7,"dim_date\ndate_sk")]: box(ax,x,5.25,label)
    box(ax,5,3.8,"fact_flight\n1 row / accepted flight ID\n+ unknown member",width=3.3,height=.88)
    for x in [1.3,5,8.7]: arrow(ax,(x,4.87),(5,4.29),"1 : many" if x==5 else None)
    box(ax,1.4,2.2,"dim_passenger\npassenger_sk",width=2.45)
    box(ax,8.6,2.2,"dim_status\nstatus_sk",width=2.2)
    box(ax,5,2.2,"fact_booking\n1 row / booking ID",width=3.3,height=.82)
    arrow(ax,(5,3.30),(5,2.67),"1 : many")
    arrow(ax,(2.69,2.2),(3.29,2.2))
    arrow(ax,(7.44,2.2),(6.72,2.2))
    arrow(ax,(8.7,4.87),(6.45,2.67),"booking date")
    box(ax,5,.64,"fact_payment\n1 row / payment ID",width=3.3,height=.82)
    arrow(ax,(5,1.74),(5,1.10),"1 : many")
    ax.text(.18,.43,"Filter path:\nflight → booking → payment",fontsize=9.2,color="#596777")
    save(fig,"data_model")


def set_styles(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    # Standard business brief preset; Code, Figure and Small are named role overrides.
    specs = {"Normal":(11,INK,0,6,1.10), "Title":(28,INK,0,10,1.05),
             "Subtitle":(12,MUTED,0,12,1.10), "Heading 1":(16,BLUE,16,8,1.10),
             "Heading 2":(13,BLUE,12,6,1.10), "Heading 3":(12,"1F4D78",8,4,1.10),
             "Caption":(9,MUTED,4,6,1.05)}
    for name,(size,color,before,after,line) in specs.items():
        style=doc.styles[name]
        style.font.name="Calibri"
        style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string(color)
        style.paragraph_format.space_before=Pt(before)
        style.paragraph_format.space_after=Pt(after)
        style.paragraph_format.line_spacing=line
        if name.startswith("Heading"): style.paragraph_format.keep_with_next=True
    for name,size,font in [("Code",9,"Consolas"),("Small",9,"Calibri"),("Figure",10,"Calibri")]:
        style=doc.styles.add_style(name,1)
        style.base_style=doc.styles["Normal"]
        style.font.name=font; style.font.size=Pt(size)
        style.paragraph_format.space_after=Pt(5)
        style.paragraph_format.line_spacing=1.05
    doc.styles["Figure"].paragraph_format.keep_with_next=True
    header=section.header.paragraphs[0]
    header.text="ASG AIRLINES  /  DATA ENGINEERING CASE STUDY"
    header.style=doc.styles["Small"]
    header.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.style=doc.styles["Small"]
    footer.add_run("Vishwas Vashishtha  |  ")
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    doc.core_properties.author="Vishwas Vashishtha"
    doc.core_properties.title="ASG Airlines: Data Engineering Case Study"
    doc.core_properties.subject="Source evidence, preprocessing, analytical model and verification"


def table(doc, headers, rows, widths):
    item=doc.add_table(rows=1,cols=len(headers))
    item.autofit=False
    item.style="Table Grid"
    props=item._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:type"),"dxa")
    props.find(qn("w:tblW")).set(qn("w:w"),"9360")
    indent=OxmlElement("w:tblInd"); indent.set(qn("w:w"),"120"); indent.set(qn("w:type"),"dxa"); props.append(indent)
    margins=OxmlElement("w:tblCellMar")
    for side,value in [("top",80),("bottom",80),("start",120),("end",120)]:
        el=OxmlElement("w:"+side);el.set(qn("w:w"),str(value));el.set(qn("w:type"),"dxa");margins.append(el)
    props.append(margins)
    for cell,text in zip(item.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(item.add_row().cells,row): cell.text=str(text)
    for column,width in zip(item.columns,widths): column.width=Inches(width/1440)
    for row_index,row in enumerate(item.rows):
        no_split=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(no_split)
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width/1440)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(3)
                paragraph.paragraph_format.line_spacing=1.05
                for run in paragraph.runs:
                    run.font.size=Pt(9.4)
                    if row_index==0:run.bold=True
            if row_index==0:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"F2F4F7");cell._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement("w:tblHeader");item.rows[0]._tr.get_or_add_trPr().append(repeat)
    return item


def build_report(root, verification_note):
    profile=json.loads((root/"reports/profile.json").read_text(encoding="utf-8"))
    manifest=json.loads((root/"reports/run_manifest.json").read_text(encoding="utf-8"))
    if manifest["status"]!="success": raise ValueError("A successful run manifest is required")
    if profile["source_sha256"]!=manifest["source_sha256"]: raise ValueError("Profile and run refer to different inputs")
    with (root/"data/gold/kpi_summary.csv").open(newline="",encoding="utf-8") as handle:
        kpi=next(csv.DictReader(handle))
    doc=Document();set_styles(doc)
    def p(text,style=None): return doc.add_paragraph(text,style)
    def h(text):doc.add_heading(text,level=1)
    def sub(text):doc.add_heading(text,level=2)
    def page(text):doc.add_page_break();h(text)
    def figure(name,caption):
        paragraph=p("", "Figure");run=paragraph.add_run();run.add_picture(str(root/"docs"/f"{name}.png"),width=Inches(6.5))
        run._r.xpath(".//wp:docPr")[0].set("descr",caption)
        p(caption,"Caption")
    money=lambda value:f"INR {Decimal(str(value)):,.2f}"
    percent=lambda value:f"{float(value)*100:.1f}%"

    p("ASG Airlines", "Title")
    p("Data engineering case study", "Subtitle")
    p("Prepared by Vishwas Vashishtha  |  Source evidence, implementation and verification", "Small")
    p("The workbook becomes a repeatable analytical dataset through a small Python and SQL pipeline. The focus is correct duration, traceable duplicate treatment, protected passenger identity and payment totals that survive joins.")
    figure("architecture","Figure 1. Workbook processing stages and analytical outputs; Power BI is the downstream consumer.")
    sub("Local result")
    p(f"The successful local run retains {int(kpi['flight_count']):,} real flights, {int(kpi['booking_count']):,} bookings and {int(kpi['payment_count']):,} payments. Valid payments total {money(kpi['gross_valid_payment_amount'])}. A source join would overstate this by {money(profile['fanout']['inflation'])}.")
    p("pandas handles the four small sheets; DuckDB executes the dimensions, facts and KPI SQL. An Azure adapter invokes the same batch. Historical cloud evidence predates the monthly extract added in this fix pass.")
    p("Reading order: source evidence; cleaning; privacy and flow; model; KPIs; execution and checks; Azure and Power BI verification.","Small")

    page("1. Source evidence")
    p("The workbook was profiled before transformation. Empty formatting padding is ignored: four unnamed flight columns and twelve blank booking rows are not business data.")
    table(doc,["Sheet","Rows","Columns","Grain in source"],[(name,str(s['rows']),str(len(s['columns'])),grain) for name,s,grain in [(n,profile['structure'][n],g) for n,g in [("flights","Flight record"),("bookings","Booking ID"),("payments","Payment ID"),("passengers","Passenger candidate")]]],[2200,1000,1100,5060])
    sub("Findings that determine the cleaning rules")
    p("Flights have 1,004 distinct IDs across 1,020 rows. Fifteen duplicate groups are exact copies. 6F250 has two conflicting rows: route origin, timestamps and duration differ. No source flight ID fails the expected pattern. There are six cities and thirty ordered routes.")
    p("Airline is null on 41 rows and literal UNKNOWN on 31. Duration is mixed: 1,019 Excel time cells and one datetime cell representing a negative serial. Departure dates span 17-20 April 2026.")
    p("Bookings span 17 April 2025 to 17 April 2026. Status counts are 320 CONFIRMED, 314 CANCELLED, 291 PENDING, 45 null and 30 INVALID. All source flight and passenger references resolve before conflicting-flight quarantine.")
    p("Amounts contain 913 floats, 9 integers, 30 INVALID strings and 48 nulls. There are 637 bookings with a payment record and 363 without. Multiple payment records per booking are real source structure.")
    sub("Corrections to the supplied verification targets")
    p("All 1,039 Aadhaar cells already contain twelve-digit strings; 114 retain a leading zero. There is no observed loss to claim. The 382 passenger rows without bookings represent 364 distinct passenger IDs. Likewise, 227 minor rows represent 215 distinct IDs.")
    p("The source has 125 date-different flight rows, including SJ192's backwards date. After its one-day correction there are 124 overnight rows before duplicate treatment; after duplicate treatment, 122 remain.")
    p("Evidence: reports/profile.md and reports/profile.json. No source PII values appear in either report.","Small")

    page("2. Cleaning and preprocessing")
    sub("Duration comes from timestamps")
    p("Calculate arrival minus departure. If negative, add exactly one calendar day to arrival and recalculate. Accept only a duration greater than zero and at most 1,440 minutes. Store was_corrected and a reason. An ordinary flight that already arrives on the next day is left as supplied.")
    table(doc,["SJ192 field","Source","Analytical result"],[('Departure','19 Apr 2026, 18:45:42','Unchanged'),('Arrival','18 Apr 2026, 23:45:42','19 Apr 2026, 23:45:42'),('Duration','-1,140 minutes','300 minutes'),('Timing flags','Backwards interval','Corrected; not overnight; not red-eye')],[2000,3450,3910])
    p("Raw duration is decoded using Excel's date epoch, preserving fractional seconds, and compared with the computed value within a one-second tolerance. It does not replace the timestamp calculation. SJ192 remains a reconciliation mismatch after repair.")
    sub("Duplicate handling preserves the intended grain")
    p("Exact flight copies collapse to one record and produce an audit issue. Both 6F250 records are quarantined because choosing one would invent an answer. Their two bookings remain and use flight_sk = -1. The result is 1,003 usable flights plus one unknown model member.")
    p("Passenger survivorship prefers fewer null business fields, then a twelve-digit Aadhaar representation, then the lexicographically smaller email. The original Excel row is the final tie-break. The 36 duplicate groups contain 75 candidate rows; 39 excess candidates collapse into 1,000 passenger IDs. The audit records row numbers and ranks without publishing emails.")
    sub("Incomplete records remain useful")
    p("Null/sentinel airline becomes UNKNOWN, with separate issue categories. Null booking status becomes UNKNOWN; INVALID remains INVALID. Both have is_valid = false. Missing and non-numeric payment amounts stay null with distinct quality labels. A missing last name uses the available first initial for the silver mask.")
    p("Implementation: src/validate.py, src/clean.py and src/pii.py. The quarantine is restricted source evidence; it is not a deletion log.","Small")

    page("3. Data flow and privacy")
    figure("data_flow","Figure 2. Structural failures stop the run; row issues determine retention, repair or quarantine.")
    sub("The analytical sharing boundary")
    p("Source and bronze retain raw PII locally. Silver applies the survivorship rule, then creates an HMAC-SHA-256 token from normalized Aadhaar using ASG_PII_PEPPER. Email, phone, date of birth, passport and emergency contact fields are removed. Gold also omits masked names and source passenger IDs.")
    p("The passenger dimension retains only the token, age, age band, gender and analytical key. A token is pseudonymous: it preserves links across records and must still be shared only with authorized reviewers. A missing or invalid identity can use a flagged, namespaced source-ID fallback; none is needed in this workbook.")
    p("The CLI requires a pepper of at least 32 characters outside Git and output files. This fix pass used one private process-only pepper for local reruns; it did not retrieve or change the cloud secret. Reuse a persistent private pepper for stable local/cloud tokens. Counts and payments do not depend on it. Git ignore rules provide no access control; current cloud permissions still need a reviewer-access test.")
    p("Evidence and limitations: docs/SECURITY.md; tests/test_pii.py and tests/test_no_pii_leak.py.","Small")

    page("4. Analytical model and join safety")
    figure("data_model","Figure 3. Unique lookup keys connect facts; payment → booking → flight is a single many-to-one path.")
    p("dim_date, dim_airline, dim_route, dim_passenger and dim_status provide the reporting context. fact_flight has one row per accepted flight ID; fact_booking has one row per booking ID; fact_payment has one row per payment ID. Unknown members retain unresolved references without creating a false route or airline.")
    table(doc,["Source join experiment","Rows","Valid payment total"],[('Bookings only','1,000','Not a payment table'),('Booking LEFT JOIN flight','1,032','Not a payment table'),('Booking LEFT JOIN payment','1,363',money(profile['fanout']['true_payment_total'])),('Naive three-way INNER JOIN','1,028',money(profile['fanout']['naive_three_way_total']))],[4700,1000,3660])
    p("The payment inner join happens to have 1,000 rows: unpaid bookings disappear while other bookings expand. Duplicate flight IDs then repeat payments. The final model checks primary keys and references and reconciles payment cents before committing its SQL transaction.")
    p("In Power BI, use single-direction relationships along the documented model path. Do not add a second direct payment-to-flight relationship.","Small")

    page("5. KPI definitions and measured results")
    p("All flight denominators exclude the synthetic unknown flight. Booking rates include every retained booking, including UNKNOWN and INVALID status. INR payment amounts remain exact decimals; displayed averages and shares are rounded only for reading.")
    table(doc,["Measure","Result","Definition / denominator"],[
        ('Real flights',kpi['flight_count'],'Accepted distinct flight IDs'),
        ('Average flight duration',f"{float(kpi['avg_duration_minutes']):.2f} min",'Sum of corrected duration / real flights'),
        ('Overnight share',percent(kpi['overnight_share']),'Different corrected arrival date / real flights'),
        ('Red-eye share',percent(kpi['red_eye_share']),'Departure 22:00-04:59 / real flights'),
        ('Cancellation rate',percent(kpi['cancellation_rate']),'CANCELLED bookings / all bookings'),
        ('Confirmation rate',percent(kpi['confirmation_rate']),'CONFIRMED bookings / all bookings'),
        ('Payment coverage',percent(kpi['payment_coverage']),'Bookings with any payment / all bookings'),
        ('Gross valid payments',money(kpi['gross_valid_payment_amount']),'Sum of valid amounts, all booking statuses'),
        ('Confirmed-booking payments',money(kpi['confirmed_booking_payment_amount']),'Valid amounts for CONFIRMED bookings'),
        ('Average paid-booking amount',money(kpi['avg_valid_payment_per_paid_booking']),'Valid amounts / distinct bookings with valid payment')
    ],[2600,2360,4400])
    sub("Required analytical views")
    p("Route traffic reports flight and booking counts separately. Airline distribution includes actual UNKNOWN-airline flights. The monthly extract contains thirteen booking periods, April 2025 through April 2026, totaling 1,000 bookings. It aggregates payments before joining bookings and includes gross valid and confirmed amounts.")
    p("The data has no scheduled-versus-actual timestamps, so delay minutes cannot be calculated. Source issues total 714 events across fifteen categories; the supplied 1,048-event target is unsupported. A row can have several issues, so events are not unique bad records. The payment-average denominator is 606 bookings with valid amounts, not all 637 with any payment.")
    p("Gross valid payments include cancelled bookings and are not net revenue or settlement. No fare, refund or boarding evidence is fabricated. SQL definitions: sql/03_kpi_views.sql; Power BI measures: powerbi/measures.dax.","Small")

    page("6. Execution, error handling and checks")
    p("Create a virtual environment, install requirements.txt and set ASG_PII_PEPPER outside the project. The authorized source workbook belongs at data/raw/UseCase_Airlines.xlsx. The .env example documents the name; it is not loaded automatically.")
    p("python -m src.pipeline", "Code")
    p("python -m pytest -q", "Code")
    p("The optional --input and --output-dir arguments support a different workbook location or output root. The call path is run_pipeline → ingest → validate/profile → clean → build_model → calculate_kpis → export_gold. The notebook imports these same modules.")
    table(doc,["Stage","Purpose","Run status"],[(s['name'],{'ingest':'Typed source and row metadata','validate':'Structure, fields and independent profile','clean':'Retention, repair, tokenization, quarantine','model':'SQL dimensions, facts and checks','kpi':'Read analytical measures','export':'Ordered CSVs and quality report'}[s['name']],s['status']) for s in manifest['stages']],[1500,6350,1510])
    sub("Visible failures and complete outputs")
    p("A per-output lock prevents concurrent local writers. Work happens in an isolated run directory. Only a completed run replaces published outputs; the prior output is restored if publication fails. Stage logs and failed-run manifests expose the failed stage and error type without printing source values or secrets.")
    p("run_manifest.json records the source hash, run ID, dependency versions, stage timings, source/silver/quarantine/gold counts and CSV hashes. A rerun uses a new run ID; stable analytical results are compared independently of timestamps and timings.")
    sub("Verification evidence")
    p(verification_note)
    p("Business checks cover SJ192, legitimate next-day arrival, distinct red-eye/overnight flags, exact and conflicting flight duplicates, preserved unknown-flight bookings, deterministic survivorship, normalized HMAC, missing-versus-invalid amounts, unique keys, source payment reconciliation and analytical PII leakage.")
    p(f"Successful local run: {manifest['run_id']}","Small")

    page("7. Azure and Power BI verification limits")
    p("Local execution and historical cloud records are separate evidence. The existing deployment manifest records an East Asia Container Apps Job, ADLS Gen2, managed identity and Key Vault. No cloud calls were made in this fix pass, and the new fourteen-extract output has not been verified there.")
    sub("Historical Azure evidence")
    p("The first saved cloud run records matching KPIs and thirteen matching gold hashes. The rerun checks blob counts; it does not establish that prior contents are immutable. The saved failure check lacks a job-exit assertion. The permissions script tests anonymous access but supplies a hard-coded role matrix, so it does not prove reviewer RBAC isolation.")
    p("The adapter uploads under run-specific paths and uses an ETag for the latest pointer. Its upload check compares sizes, not downloaded remote hashes; overwrite is enabled. The deployment downloads a mutable source package and installs additional dependencies at runtime. These limits must be addressed or stated when refreshing cloud evidence.")
    sub("Power BI project status")
    p("The PBIP source is retained with twelve table definitions, twenty-six measures, eight relationships and four page definitions. The generated PBIX/PBIT and five rendered images were removed because they were not evidence of a working Desktop report. No replacement screenshots were generated.")
    p("The semantic model still references missing CSV columns, including dim_route.origin and several is_unknown/is_valid fields. Its imports use local File.Contents; the Azure URL and run parameters do not drive them. Correct these mappings in Desktop, connect every table to one successful run, then refresh and reconcile the report.")
    p("Use kpi_booking_month.csv to verify all thirteen actual year-month periods. Check unknown-flight bookings, multi-payment totals, slicers and date roles. Save and reopen the real PBIX, then capture genuine page/model screenshots. powerbi/VERIFICATION.md lists the remaining checks.")
    sub("Submission evidence")
    p("FIX_REPORT.md records fixes, independent figures and remaining findings. Keep raw data and credentials outside Git. Historical cloud files have been preserved without changing their original labels; their limitations are documented alongside the current local result.")
    p("Useful companion files: docs/DECISIONS.md, docs/ASSUMPTIONS.md, docs/EXECUTION_FLOW.md and docs/EXPLAINABILITY_CHECKLIST.md. The current implementation is a single-node full refresh; larger workloads require new measurements before adding infrastructure.","Small")

    output=root/"docs/ASG_Airlines_Case_Study.docx"
    doc.save(output)
    print(f"Created {output}")


def main():
    parser=argparse.ArgumentParser(description="Build source-backed case-study documentation.")
    parser.add_argument("--root",type=Path,default=ROOT)
    parser.add_argument("--diagrams-only",action="store_true")
    parser.add_argument("--verification-note",default="The six pipeline stages completed successfully. Test and independent rerun evidence must be reviewed in the repository before submission.")
    args=parser.parse_args()
    draw_diagrams(args.root/"docs")
    if not args.diagrams_only:build_report(args.root,args.verification_note)


if __name__=="__main__":main()
